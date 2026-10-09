# -*- coding: utf-8 -*-
"""
BÁO CÁO CHẤT LƯỢNG CÔNG VIỆC PHÁP CHẾ - PHÒNG KIỂM SOÁT NỘI BỘ
Chạy:  streamlit run app.py

- Chỉ ĐỌC dữ liệu nguồn (Google Drive), không ghi/sửa gì vào file gốc.
- 1 trang duy nhất, bộ lọc bên trái.
- Nhận xét bổ sung được lưu ở file nhan_xet_bo_sung.json cạnh app.py (không đụng tới dữ liệu nguồn).
"""
import base64
import html
import io
import json
import math
import re
import textwrap
import unicodedata
from datetime import date
from pathlib import Path

import numpy as np
import pandas as pd
import plotly.graph_objects as go
import streamlit as st

# ----------------------------------------------------------------------------
# CẤU HÌNH
# ----------------------------------------------------------------------------
DEFAULT_SHEET_ID = "131BlpOXvjyDMjxaQRkEgCDuQkWF3DaWN"


def source_urls() -> list:
    """Thử lần lượt ID trong Secrets (nếu có) rồi ID mặc định; mỗi ID thử link export Sheet và link tải Drive."""
    ids = []
    try:
        sec = str(st.secrets.get("SHEET_ID", "") or "").strip()
        if sec:
            ids.append(sec)
    except Exception:
        pass
    if DEFAULT_SHEET_ID and DEFAULT_SHEET_ID not in ids:
        ids.append(DEFAULT_SHEET_ID)
    urls = []
    for sid in ids:
        urls.append(f"https://docs.google.com/spreadsheets/d/{sid}/export?format=xlsx")   # Google Sheet gốc
        urls.append(f"https://drive.google.com/uc?export=download&id={sid}")             # file .xlsx lưu trên Drive
    return urls


try:  # giờ Việt Nam (server Streamlit Cloud chạy giờ UTC)
    from zoneinfo import ZoneInfo
    from datetime import datetime
    TODAY = datetime.now(ZoneInfo("Asia/Ho_Chi_Minh")).date()
except Exception:
    TODAY = date.today()
LOGO_URL = (
    "https://scontent.fhan12-1.fna.fbcdn.net/v/t39.30808-6/584260385_837902692321611_5716056316"
    "288511031_n.jpg?stp=dst-jpg_tt6&cstp=mx2048x2048&ctp=s2048x2048&_nc_cat=101&_nc_map=urlgen"
    "_bucketless&ccb=1-7&_nc_sid=6ee11a&_nc_ohc=g8KidzZ6L5QQ7kNvwGNfgL_&_nc_oc=AdrYQ6ZR6jg9BxjW"
    "Klg8E2vu_5NfNZtnsopCByiqkhGV9pWHV4IQ23RzS0P7nOedJbw&_nc_zt=23&_nc_ht=scontent.fhan12-1.fna"
    "&_nc_gid=hjpRNR4cui61LBT0GMwqeg&_nc_ss=7b2a8&oh=00_AQNXYuUUwXpDVwVWQDuX_MDxWv4_JbucOQ6OrNY"
    "rKVOZRA&oe=6ACD8ECB")
NOTES_FILE = Path(__file__).with_name("nhan_xet_bo_sung.json")

ALIASES = {
    "detail":   ["chi tiet yeu cau", "chi tiet", "noi dung yeu cau", "noi dung"],
    "status":   ["trang thai thuc hien", "trang thai"],
    "received": ["ngay nhan yeu cau", "ngay tiep nhan", "ngay nhan"],
    "done":     ["ngay hoan thanh", "ngay hoan tat"],
    "owner":    ["can bo phu trach", "nguoi thuc hien", "nguoi phu trach"],
    "note":     ["ghi chu"],
    "category": ["hang muc", "nhom cong viec", "loai cong viec"],
    "unit":     ["don vi yeu cau", "don vi", "phong ban yeu cau"],
    "quarter":  ["quy"],
    "year":     ["nam"],
}
FIELD_LABEL = {
    "detail": "Chi tiết yêu cầu", "status": "Trạng thái", "received": "Ngày nhận yêu cầu",
    "done": "Ngày hoàn thành", "owner": "Cán bộ phụ trách", "note": "Ghi chú",
    "category": "Hạng mục", "unit": "Đơn vị yêu cầu", "quarter": "Quý (nếu có sẵn)", "year": "Năm (nếu có sẵn)",
}
# Chuẩn hóa trạng thái (so khớp sau khi bỏ dấu, KHÔNG dùng "chứa" để tránh "Chưa hoàn thành" bị tính là hoàn thành)
STATUS_CANON = {
    "hoan thanh": "Hoàn thành", "da hoan thanh": "Hoàn thành", "hoan tat": "Hoàn thành",
    "dang thuc hien": "Đang thực hiện", "dang xu ly": "Đang thực hiện", "dang trien khai": "Đang thực hiện",
    "tam dung": "Tạm dừng", "tam hoan": "Tạm dừng",
}
ST_DONE, ST_OPEN, ST_PAUSE = "Hoàn thành", "Đang thực hiện", "Tạm dừng"

TEAL, TEAL_DARK, TEAL_DEEP = "#0FA89B", "#0B7F75", "#075E57"
INK, MUTED, LINE, BG = "#16322F", "#6B8581", "#E4EEEC", "#F3F8F7"
BLUE, ORANGE, NAVY, GREY = "#2D9CDB", "#F2994A", "#27348B", "#9AA5A4"
STATUS_COLORS = {ST_DONE: BLUE, ST_OPEN: ORANGE, ST_PAUSE: NAVY, "Chưa xác định": GREY}
PERSON_COLORS = [BLUE, "#F2A07B", NAVY, "#2BB5A6", "#8E6BBF", "#E0B400", GREY]
UNIT_COLORS = [BLUE, NAVY, ORANGE, "#7B1FA2", "#E91E8C", "#8E6BBF", "#E0B400", "#2BB5A6", GREY]
FONT = "Be Vietnam Pro, Segoe UI, Roboto, Arial, sans-serif"
ROMAN = {1: "I", 2: "II", 3: "III", 4: "IV"}

st.set_page_config(page_title="Báo cáo chất lượng công việc Pháp chế", page_icon="⚖️",
                   layout="wide", initial_sidebar_state="expanded")

CSS = """
<style>
#MainMenu, footer, [data-testid="stMainMenu"], [data-testid="stAppDeployButton"], [data-testid="stToolbarActions"], [data-testid="stDecoration"], [data-testid="stStatusWidget"]{visibility:hidden;display:none;}
[data-testid="stSidebarCollapsedControl"], [data-testid="stExpandSidebarButton"], [data-testid="stSidebarCollapseButton"]{visibility:visible !important;display:flex !important;}

@import url('https://fonts.googleapis.com/css2?family=Be+Vietnam+Pro:wght@400;500;600;700;800&display=swap');
.stApp {font-family:'Be Vietnam Pro','Segoe UI',Roboto,Arial,sans-serif; background:__BG__;}
header[data-testid="stHeader"] {background:transparent;}
.block-container {padding-top:3.8rem; padding-bottom:3rem; max-width:1560px;}
section[data-testid="stSidebar"] {background:#fff;}

/* ---------- Header ---------- */
.hero {display:flex; align-items:center; gap:22px; padding:22px 28px; border-radius:16px; color:#fff;
       background:linear-gradient(120deg,__TEAL_DEEP__ 0%,__TEAL_DARK__ 48%,__TEAL__ 100%);
       box-shadow:0 8px 22px rgba(11,127,117,.25);}
.hero > img {display:block; height:104px; width:auto; max-width:380px; object-fit:contain; border-radius:12px;
             box-shadow:0 2px 8px rgba(0,0,0,.18);}
.hero .logo-tile {background:#fff; border-radius:10px; padding:4px 8px; display:flex; align-items:center;}
.hero .logo-tile img {display:block; height:64px; width:auto; max-width:280px; object-fit:contain;}
.hero .eyebrow {letter-spacing:.2em; font-size:.7rem; font-weight:700; color:#C4F5EE; text-transform:uppercase;}
.hero .title {margin:5px 0 4px 0; font-size:1.5rem; line-height:1.3; font-weight:800; color:#fff;}
.hero .sub {font-size:.86rem; opacity:.93;}
.hero .pill {margin-left:auto; white-space:nowrap; border:1px solid rgba(255,255,255,.55); border-radius:999px;
             padding:8px 18px; font-size:.85rem; font-weight:600; background:rgba(255,255,255,.12);}

/* ---------- Section ---------- */
.sec {margin:38px 0 14px 0; padding-left:14px; border-left:5px solid __TEAL__;}
.sec .eb {color:__TEAL__; font-weight:800; letter-spacing:.16em; font-size:.72rem;}
.sec .tt {color:__INK__; font-weight:800; font-size:1.32rem; margin-top:1px; line-height:1.3;}
.sec .st {color:__MUTED__; font-size:.84rem; margin-top:1px;}

/* ---------- KPI ---------- */
.kpi {background:#fff; border:1px solid __LINE__; border-top:4px solid __TEAL__; border-radius:12px;
      padding:12px 16px 12px 16px; min-height:138px; box-shadow:0 1px 4px rgba(15,60,55,.06);}
.kpi.ok {border-top-color:__BLUE__;} .kpi.warn {border-top-color:__ORANGE__;} .kpi.pause {border-top-color:__NAVY__;}
.kpi .kl {color:__MUTED__; font-size:.72rem; font-weight:700; letter-spacing:.07em; text-transform:uppercase;}
.kpi .kv {color:__INK__; font-size:2rem; font-weight:800; line-height:1.15; margin:3px 0 6px 0;}
.kpi .kv small {font-size:.9rem; font-weight:600; color:__MUTED__;}
.kpi .kr {display:flex; justify-content:space-between; font-size:.78rem; color:__MUTED__; margin-bottom:6px;}
.kpi .kr b {color:__INK__;}
.kd {display:inline-block; font-size:.74rem; font-weight:700; padding:3px 9px; border-radius:7px;}
.kd.good {background:#DFF6EE; color:#0B7F5B;} .kd.bad {background:#FDE7DD; color:#C2501F;}
.kd.flat {background:#EEF3F2; color:#5E7773;}

/* ---------- Card ---------- */
div[class*="st-key-card_"] {background:#fff; border:1px solid __LINE__; border-radius:14px;
      box-shadow:0 1px 4px rgba(15,60,55,.06); padding:12px 16px 18px 16px;}
.ct {font-weight:800; color:__INK__; font-size:1rem;}
.cs {color:__MUTED__; font-size:.78rem; margin-bottom:4px;}

/* ---------- Bảng tự nhiên ---------- */
table.nt {width:100%; border-collapse:collapse; border:0; font-size:.87rem; color:__INK__; font-variant-numeric:tabular-nums;}
table.nt th, table.nt td {border:0; background:transparent;}
table.nt th {font-size:.72rem; font-weight:700; color:__MUTED__; text-align:right; padding:8px 10px;
             border-bottom:2px solid __TEAL__; vertical-align:bottom;}
table.nt thead tr.g1 th {border-bottom:1px solid #DCE8E5; text-align:center; color:__INK__; font-size:.8rem;}
table.nt th:first-child, table.nt td:first-child {text-align:left;}
table.nt td {padding:8px 10px; text-align:right; border-bottom:1px solid #EDF3F2;}
table.nt tbody tr:hover td {background:#F2FAF8;}
table.nt tr.tot td {font-weight:800; border-top:2px solid #C9DAD7; border-bottom:0;}
table.nt td.dim {color:#A5B8B4;}
table.nt td.low {color:#C2501F; font-weight:700;}
table.nt th.grp, table.nt td.grp {border-left:1px solid #E6EFED;}
.rb {display:flex; align-items:center; gap:8px; justify-content:flex-end;}
.rb .bar {width:64px; height:7px; background:#E8F1EF; border-radius:4px; overflow:hidden;}
.rb .bar i {display:block; height:100%; background:__TEAL__; border-radius:4px;}
.rb.warn .bar i {background:__ORANGE__;}
.dot {display:inline-block; width:9px; height:9px; border-radius:50%; margin-right:7px;}

/* ---------- Kết luận ---------- */
.concl {font-size:.92rem; line-height:1.6; color:__INK__;}
.concl h4 {margin:0 0 6px 0; font-size:1.02rem; font-weight:800; padding:0;}
.concl ul {padding-left:18px; margin:0;}
.concl li {margin-bottom:4px;}
.concl .xt {margin:10px 0 4px 0; font-weight:800; color:__TEAL_DARK__; font-size:.82rem; letter-spacing:.05em;}
.legend-row {display:flex; justify-content:space-between; font-size:.85rem; padding:3px 2px; color:__INK__;}
.legend-row span.n {color:__MUTED__;}
.fn {color:__MUTED__; font-size:.78rem; line-height:1.65; margin:10px 2px 0 2px;}
.fn b {color:__INK__;}
</style>
"""
for k, v in {"__BG__": BG, "__TEAL_DEEP__": TEAL_DEEP, "__TEAL_DARK__": TEAL_DARK, "__TEAL__": TEAL, "__INK__": INK,
             "__MUTED__": MUTED, "__LINE__": LINE, "__BLUE__": BLUE, "__ORANGE__": ORANGE, "__NAVY__": NAVY}.items():
    CSS = CSS.replace(k, v)
st.markdown(CSS, unsafe_allow_html=True)


# ----------------------------------------------------------------------------
# TIỆN ÍCH
# ----------------------------------------------------------------------------
def norm(s) -> str:
    s = unicodedata.normalize("NFD", str(s))
    s = "".join(c for c in s if unicodedata.category(c) != "Mn").replace("đ", "d").replace("Đ", "D")
    return re.sub(r"\s+", " ", s).strip().lower()


def vn(x, d=0) -> str:
    """Định dạng số kiểu Việt Nam: 1.234,5"""
    if x is None or (isinstance(x, float) and np.isnan(x)) or pd.isna(x):
        return "–"
    s = f"{x:,.{d}f}"
    return s.replace(",", "§").replace(".", ",").replace("§", ".")


def rate_txt(done, total, d=1) -> str:
    """Tỷ lệ % làm tròn XUỐNG; chỉ hiện 100% khi toàn bộ đã hoàn thành (tránh 99,8% bị làm tròn thành 100%)."""
    if not total:
        return "–"
    if done >= total:
        return "100%"
    v = math.floor(done / total * 10 ** (d + 2) + 1e-9) / 10 ** d
    return f"{vn(v, d)}%"


def share_txt(a, b, d=1) -> str:
    """Tỷ trọng % thông thường; nếu có phát sinh nhưng quá nhỏ thì hiện '<0,1%' thay vì 0%."""
    if not b:
        return "–"
    if a <= 0:
        return "0%"
    if a >= b:
        return "100%"
    v = a / b * 100
    if round(v, d) <= 0:
        return f"<{vn(10 ** -d, d)}%"
    if round(v, d) >= 100:
        return rate_txt(a, b, d)
    return f"{vn(v, d)}%"


def auto_map(columns) -> dict:
    normed = {c: norm(c) for c in columns}
    mapping, used = {}, set()
    for field, names in ALIASES.items():
        found = None
        for n in names:
            for c, nc in normed.items():
                if nc == n and c not in used:
                    found = c
                    break
            if found:
                break
        if not found and field not in ("quarter", "year"):
            for n in names:
                for c, nc in normed.items():
                    if n in nc and c not in used:
                        found = c
                        break
                if found:
                    break
        if found:
            mapping[field] = found
            used.add(found)
    return mapping


def parse_date(s: pd.Series) -> pd.Series:
    if pd.api.types.is_datetime64_any_dtype(s):
        return s
    out = pd.to_datetime(s, errors="coerce", format="%m/%d/%Y")
    miss = out.isna() & s.notna()
    if miss.any():
        out.loc[miss] = pd.to_datetime(s[miss], errors="coerce", dayfirst=True)
    return out


def canon_status(s) -> str:
    n = norm(s)
    if not n or n in ("nan", "none", "nat"):
        return "Chưa xác định"
    return STATUS_CANON.get(n, str(s).strip())


def canon_labels(s: pd.Series, empty: str) -> pd.Series:
    """Gộp các cách viết khác nhau (hoa/thường, dấu cách, NFC/NFD) về cùng 1 nhãn."""
    s = s.fillna("").astype(str).str.strip().str.replace(r"\s+", " ", regex=True)
    s = s.replace({"nan": "", "None": "", "NaT": ""})
    key = s.map(norm)
    def _pick(x):
        vc = x.value_counts()
        tops = vc[vc == vc.max()].index          # cùng số lần xuất hiện -> ưu tiên cách viết có nhiều chữ hoa (đúng chính tả)
        return max(tops, key=lambda v: sum(ch.isupper() for ch in v))
    best = s.groupby(key).agg(_pick)
    out = key.map(best).astype(object)
    return out.where(key != "", empty)


@st.cache_data(ttl=600, show_spinner="Đang tải dữ liệu nguồn...")
def fetch_source():
    import requests
    for url in source_urls():
        try:
            r = requests.get(url, timeout=40, allow_redirects=True)
            if r.status_code == 200 and r.content[:2] == b"PK":
                return r.content
        except Exception:
            continue
    raise RuntimeError("Không tải được dữ liệu nguồn")


def _process_logo(raw: bytes) -> dict:
    """Xử lý ảnh logo: cắt viền trong suốt, thu nhỏ, mã hóa base64; nhận biết logo sáng/nền đặc để không cần ô trắng."""
    from PIL import Image
    img = Image.open(io.BytesIO(raw)).convert("RGBA")
    a = np.asarray(img)
    has_alpha = a[..., 3].min() < 250
    if has_alpha:
        bbox = img.getchannel("A").point(lambda v: 255 if v > 20 else 0).getbbox()
        if bbox:
            img = img.crop(bbox)
            a = np.asarray(img)
    m = a[..., 3] > 128
    lum = float((0.299 * a[..., 0] + 0.587 * a[..., 1] + 0.114 * a[..., 2])[m].mean()) if m.any() else 0.0
    img.thumbnail((700, 220))
    buf = io.BytesIO()
    img.save(buf, "PNG")
    return {"b64": base64.b64encode(buf.getvalue()).decode(), "plain": bool((not has_alpha) or lum > 200), "src": None}


@st.cache_data(ttl=3600, show_spinner=False)
def load_logo() -> dict:
    # 1) File logo nằm trong repo (logo.png / logo.jpg ...) - KHUYÊN DÙNG: link ảnh Facebook có chữ ký, hết hạn sau vài ngày
    here = Path(__file__).parent
    for name in ["logo.png", "logo.jpg", "logo.jpeg", "logo.webp"]:
        fp = here / name
        if fp.exists():
            try:
                return _process_logo(fp.read_bytes())
            except Exception:
                pass
    # 2) Tải từ link LOGO_URL
    import requests
    from urllib.parse import parse_qs, unquote, urlparse
    cands = [LOGO_URL]
    try:
        direct = parse_qs(urlparse(LOGO_URL).query).get("url", [None])[0]
        if direct:
            cands.append(unquote(direct))
    except Exception:
        pass
    for u in cands:
        try:
            r = requests.get(u, timeout=20, headers={"User-Agent": "Mozilla/5.0"})
            if r.status_code == 200 and r.headers.get("content-type", "").startswith("image"):
                return _process_logo(r.content)
        except Exception:
            continue
    return {"b64": None, "plain": True, "src": LOGO_URL}


def best_header(raw: pd.DataFrame):
    best, best_i = -1, 0
    for i in range(min(15, len(raw))):
        score = len(auto_map([str(x) for x in raw.iloc[i].tolist() if pd.notna(x)]))
        if score > best:
            best, best_i = score, i
    df = raw.iloc[best_i + 1:].copy()
    df.columns = [str(c).strip() if pd.notna(c) else f"cot_{j}" for j, c in enumerate(raw.iloc[best_i])]
    return df.dropna(how="all").reset_index(drop=True), best


def prepare(df: pd.DataFrame, mp: dict) -> pd.DataFrame:
    out = pd.DataFrame(index=df.index)
    out["detail"] = df[mp["detail"]].fillna("").astype(str).str.strip() if "detail" in mp else ""
    out["note"] = df[mp["note"]].fillna("").astype(str).str.strip() if "note" in mp else ""
    out["status"] = df[mp["status"]].map(canon_status) if "status" in mp else "Chưa xác định"
    out["owner"] = canon_labels(df[mp["owner"]], "(Chưa phân công)") if "owner" in mp else "(Chưa phân công)"
    out["category"] = canon_labels(df[mp["category"]], "Khác") if "category" in mp else "Khác"
    out["unit"] = canon_labels(df[mp["unit"]], "(Không rõ)") if "unit" in mp else "(Không rõ)"
    out["received"] = parse_date(df[mp["received"]])
    out["done"] = parse_date(df[mp["done"]]) if "done" in mp else pd.NaT

    yr = pd.to_numeric(df[mp["year"]], errors="coerce") if "year" in mp else pd.Series(np.nan, index=out.index)
    out["year"] = yr.fillna(out["received"].dt.year)
    if "quarter" in mp:
        qq = pd.to_numeric(df[mp["quarter"]].astype(str).str.extract(r"(\d)")[0], errors="coerce")
    else:
        qq = pd.Series(np.nan, index=out.index)
    out["q"] = qq.fillna(out["received"].dt.quarter)

    out = out[(out["detail"] != "") | out["received"].notna()].copy()
    out["year"] = out["year"].astype("Int64")
    out["q"] = out["q"].astype("Int64")
    out["quarter"] = "Q" + out["q"].astype(str)
    out["ym"] = out["received"].dt.to_period("M").astype(str)

    out["is_done"] = out["status"].eq(ST_DONE)
    # số ngày LỊCH giữa 2 ngày (bỏ giờ phút nếu có): cùng ngày = 0, hôm sau = 1
    raw_days = (out["done"].dt.normalize() - out["received"].dt.normalize()).dt.days
    # Thời gian xử lý: CHỈ tính case đã hoàn thành, có đủ 2 ngày hợp lệ
    out["days"] = raw_days.where(out["is_done"] & out["done"].notna() & (raw_days >= 0))
    # Cờ kiểm tra chất lượng dữ liệu
    out["dq_done_nodate"] = out["is_done"] & out["done"].isna()
    out["dq_date_notdone"] = ~out["is_done"] & out["done"].notna()
    out["dq_negative"] = raw_days < 0
    out["dq_other_status"] = ~out["status"].isin([ST_DONE, ST_OPEN, ST_PAUSE])
    return out.reset_index(drop=True)


def to_excel(df: pd.DataFrame) -> bytes:
    buf = io.BytesIO()
    x = df.copy()
    for c in ["Ngày nhận yêu cầu", "Ngày hoàn thành"]:
        if c in x:
            x[c] = pd.to_datetime(x[c]).dt.date
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        x.to_excel(w, index=False, sheet_name="Chi tiết công việc")
        ws = w.sheets["Chi tiết công việc"]
        from openpyxl.styles import Alignment, Font, PatternFill
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="0FA89B")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for i, col in enumerate(x.columns, 1):
            width = min(max(len(str(col)), x[col].astype(str).str.len().quantile(.95) if len(x) else 10) + 2, 60)
            ws.column_dimensions[ws.cell(1, i).column_letter].width = width
            if "Ngày" in col and "Số" not in col:
                for r in range(2, len(x) + 2):
                    ws.cell(r, i).number_format = "dd/mm/yyyy"
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
    return buf.getvalue()


# ---------- thống kê ----------
def stats(d: pd.DataFrame) -> dict:
    n = len(d)
    done = int(d["is_done"].sum())
    op = int(d["status"].eq(ST_OPEN).sum())
    pa = int(d["status"].eq(ST_PAUSE).sum())
    return dict(n=n, done=done, op=op, pa=pa, oth=n - done - op - pa, cat=d["category"].nunique())


def proc_stats(d: pd.DataFrame) -> dict:
    dv = d[d["days"].notna()]
    n = len(dv)
    return dict(n=n, avg=dv["days"].mean() if n else np.nan,
                same=int((dv["days"] == 0).sum()), slow=int((dv["days"] > 7).sum()))


# ---------- thành phần giao diện ----------
def delta_pill(cur, prev, kind="neutral", mode="pct", unit="", partial=False) -> str:
    if prev is None:
        return ""
    if partial:
        return "<span class='kd flat'>Quý chưa kết thúc</span>"
    if cur is None or prev is None or pd.isna(cur) or pd.isna(prev):
        return "<span class='kd flat'>–</span>"
    diff = cur - prev
    if abs(diff) < 1e-9:
        return "<span class='kd flat'>● không đổi</span>"
    arrow = "▲" if diff > 0 else "▼"
    sign = "+" if diff > 0 else ""
    if mode == "pts":
        txt = f"{arrow} {sign}{vn(diff, 1)} điểm %"
    elif mode == "pct":
        txt = f"{arrow} {sign}{vn(diff / prev * 100, 1)}%" if prev else f"{arrow} {sign}{vn(diff)}"
    else:
        txt = f"{arrow} {sign}{vn(diff, 1 if mode == 'dec' else 0)}{unit}"
    if kind == "up_good":
        cls = "good" if diff > 0 else "bad"
    elif kind == "up_bad":
        cls = "bad" if diff > 0 else "good"
    else:
        cls = "flat"
    return f"<span class='kd {cls}'>{txt}</span>"


def kpi(label, value, prev_txt=None, pill="", cls="", note="", prev_label="Kỳ trước") -> str:
    if prev_txt is not None:
        foot = f"<div class='kr'><span>{prev_label}</span><b>{prev_txt}</b></div><div>{pill}</div>"
    else:
        foot = f"<div class='kr'><span>{note}</span></div>"
    return f"<div class='kpi {cls}'><div class='kl'>{label}</div><div class='kv'>{value}</div>{foot}</div>"


def section(eyebrow, title, sub=""):
    st.markdown(f"<div class='sec'><div class='eb'>{eyebrow}</div><div class='tt'>{title}</div>"
                f"<div class='st'>{sub}</div></div>", unsafe_allow_html=True)


def ct(title, sub=None):
    st.markdown(f"<div class='ct'>{title}</div>" + (f"<div class='cs'>{sub}</div>" if sub else ""),
                unsafe_allow_html=True)


def card(key):
    return st.container(key=f"card_{key}")


def style_fig(fig, h=300, legend=False, top=8):
    fig.update_layout(
        height=h, margin=dict(l=6, r=6, t=34 if legend else top, b=6), showlegend=legend,
        paper_bgcolor="rgba(0,0,0,0)", plot_bgcolor="rgba(0,0,0,0)", separators=",.",
        font=dict(family=FONT, size=12, color=INK), hoverlabel=dict(font_family=FONT),
        legend=dict(orientation="h", yanchor="bottom", y=1.0, xanchor="left", x=0, title_text="", font=dict(size=11),
                    traceorder="normal"),
        uniformtext=dict(minsize=9, mode="hide"),
    )
    fig.update_traces(textangle=0, selector=dict(type="bar"))
    fig.update_xaxes(showgrid=False, automargin=True, linecolor="#D5E3E0", ticks="")
    fig.update_yaxes(gridcolor="#EDF3F2", zeroline=False, automargin=True, ticks="")
    return fig


def show(fig):
    st.plotly_chart(fig, width="stretch", config={"displayModeBar": False})


def rate_cell(done, total, cls_td="") -> str:
    if not total:
        return "<td></td>"
    v = done / total * 100
    warn = "warn" if v < 90 else ""
    return (f"<td class='{cls_td}'><div class='rb {warn}'><div class='bar'><i style='width:{v:.1f}%'></i></div>"
            f"<span>{rate_txt(done, total)}</span></div></td>")


def num_cell(v, cls="") -> str:
    return f"<td class='{cls}'>{vn(v)}</td>" if v else f"<td class='dim {cls}'>–</td>"


def pivot_html(f: pd.DataFrame, people: list, rows: list, row_label: str) -> str:
    g = f.groupby(["category", "owner"]).agg(n=("detail", "size"), d=("is_done", "sum"))
    h = f"<table class='nt'><thead><tr class='g1'><th rowspan='2' style='border-bottom:2px solid {TEAL};text-align:left'>{row_label}</th>"
    for p in people:
        h += f"<th colspan='2' class='grp'>{html.escape(p)}</th>"
    h += "<th colspan='2' class='grp'>Tổng</th></tr><tr>"
    for _ in people + ["Tổng"]:
        h += "<th class='grp'>Yêu cầu</th><th>Hoàn thành</th>"
    h += "</tr></thead><tbody>"
    for r in rows:
        h += f"<tr><td>{html.escape(r)}</td>"
        tn = td = 0
        for p in people:
            n, d = (int(g.at[(r, p), "n"]), int(g.at[(r, p), "d"])) if (r, p) in g.index else (0, 0)
            tn += n
            td += d
            low = "low" if n and d < n else ""
            h += (f"<td class='grp'>{vn(n) if n else ''}</td>"
                  f"<td class='{low}'>{rate_txt(d, n) if n else ''}</td>")
        low = "low" if td < tn else ""
        h += f"<td class='grp'>{vn(tn)}</td><td class='{low}'>{rate_txt(td, tn)}</td></tr>"
    h += "<tr class='tot'><td>Tổng</td>"
    gn = gd = 0
    for p in people:
        sub = f[f["owner"] == p]
        n, d = len(sub), int(sub["is_done"].sum())
        gn += n
        gd += d
        h += f"<td class='grp'>{vn(n)}</td><td>{rate_txt(d, n)}</td>"
    h += f"<td class='grp'>{vn(gn)}</td><td>{rate_txt(gd, gn)}</td></tr></tbody></table>"
    return h


def load_notes() -> dict:
    try:
        return json.loads(NOTES_FILE.read_text(encoding="utf-8"))
    except Exception:
        return {}


def save_note(period_key: str, widget_key: str):
    notes = load_notes()
    txt = st.session_state.get(widget_key, "").strip()
    if txt:
        notes[period_key] = txt
    else:
        notes.pop(period_key, None)
    try:
        NOTES_FILE.write_text(json.dumps(notes, ensure_ascii=False, indent=2), encoding="utf-8")
    except Exception:
        pass


# ----------------------------------------------------------------------------
# NẠP DỮ LIỆU
# ----------------------------------------------------------------------------
try:
    content = fetch_source()
except Exception:
    content = None
with st.sidebar:
    st.markdown("### 🔎 Bộ lọc")
    if st.button("🔄 Làm mới dữ liệu", width="stretch"):
        fetch_source.clear()
        load_logo.clear()
        st.rerun()

if content is None:
    st.warning("Chưa tự tải được dữ liệu từ Google Drive (file cần để chế độ *Anyone with the link – Viewer*). "
               "Bạn có thể tải file về (File ▸ Download ▸ .xlsx) rồi upload vào đây — dữ liệu gốc không bị thay đổi.")
    uploaded = st.file_uploader("Upload file nguồn (.xlsx)", type=["xlsx"])
    if uploaded is None:
        st.stop()
    content = uploaded.getvalue()

sheets = pd.read_excel(io.BytesIO(content), sheet_name=None, header=None, dtype=object)
scored = {}
for name, raw in sheets.items():
    try:
        scored[name] = best_header(raw)
    except Exception:
        continue
if not scored:
    st.error("Không đọc được sheet nào trong file.")
    st.stop()
default_sheet = max(scored, key=lambda k: (scored[k][1], len(scored[k][0])))

with st.sidebar.expander("⚙️ Nguồn dữ liệu / ghép cột", expanded=False):
    sheet_name = st.selectbox("Sheet", list(scored), index=list(scored).index(default_sheet))
    df_raw, _ = scored[sheet_name]
    mp = auto_map(df_raw.columns)
    cols = ["(không có)"] + list(df_raw.columns)
    for fld, lab in FIELD_LABEL.items():
        cur = mp.get(fld)
        sel = st.selectbox(lab, cols, index=cols.index(cur) if cur in cols else 0, key=f"map_{fld}")
        if sel == "(không có)":
            mp.pop(fld, None)
        else:
            mp[fld] = sel

if "received" not in mp:
    st.error("Không tìm thấy cột **Ngày nhận yêu cầu**. Hãy chọn cột trong mục ⚙️ Nguồn dữ liệu / ghép cột ở thanh bên trái.")
    st.stop()

data = prepare(df_raw, mp)
if data.empty:
    st.error("Không có dòng dữ liệu hợp lệ.")
    st.stop()
n_no_received = int(data["received"].isna().sum())
last_date = data["received"].max()

# ----------------------------------------------------------------------------
# BỘ LỌC BÊN TRÁI
# ----------------------------------------------------------------------------
years = sorted(data["year"].dropna().unique().tolist())
latest_year = years[-1]
q_latest = int(data.loc[data["year"] == latest_year, "q"].max())

# Quý mặc định: quý gần nhất đã kết thúc đủ 3 tháng (có dữ liệu); nếu chưa có thì lấy quý mới nhất
_yq = data[["year", "q"]].dropna().drop_duplicates()
_ended = [(int(y), int(q)) for y, q in zip(_yq["year"], _yq["q"])
          if pd.Timestamp(TODAY) > pd.Timestamp(int(y), 3 * int(q), 1) + pd.offsets.MonthEnd(0)]
def_year, def_q_num = max(_ended) if _ended else (int(latest_year), q_latest)

with st.sidebar:
    sel_year = st.selectbox("Năm", years, index=years.index(def_year) if def_year in years else len(years) - 1)
    quarters = sorted(data.loc[data["year"] == sel_year, "quarter"].dropna().unique().tolist())
    default_q = [f"Q{def_q_num}"] if (sel_year == def_year and f"Q{def_q_num}" in quarters) else quarters
    sel_q = st.multiselect("Quý", quarters, default=default_q)
    sel_unit = st.multiselect("Đơn vị yêu cầu", sorted(data["unit"].unique()))
    sel_status = st.multiselect("Trạng thái", sorted(data["status"].unique()))
    sel_owner = st.multiselect("Người thực hiện", sorted(data["owner"].unique()))
    sel_cat = st.multiselect("Hạng mục", sorted(data["category"].unique()))
    keyword = st.text_input("Tìm theo chi tiết yêu cầu", placeholder="VD: hợp đồng, BBNT...")
    st.caption(f"Dữ liệu: {len(data):,} yêu cầu · tự cập nhật tối đa mỗi 10 phút")


def apply(df, year, quarters_, use_q=True):
    m = df["year"] == year
    if use_q and quarters_:
        m &= df["quarter"].isin(quarters_)
    if sel_unit:
        m &= df["unit"].isin(sel_unit)
    if sel_status:
        m &= df["status"].isin(sel_status)
    if sel_owner:
        m &= df["owner"].isin(sel_owner)
    if sel_cat:
        m &= df["category"].isin(sel_cat)
    if keyword:
        m &= df["detail"].str.contains(re.escape(keyword), case=False, na=False)
    return df[m]


f = apply(data, sel_year, sel_q)
f_allq = apply(data, sel_year, sel_q, use_q=False)

# kỳ trước (chỉ khi chọn đúng 1 quý)
prev_df, prev_label, prev_key = None, "Kỳ trước", None
if len(sel_q) == 1:
    qn = int(sel_q[0][1:])
    py, pq = (sel_year, qn - 1) if qn > 1 else (sel_year - 1, 4)
    pdf = apply(data, py, [f"Q{pq}"])
    if len(pdf):
        prev_df, prev_label, prev_key = pdf, f"Q{pq}/{py}", (py, [f"Q{pq}"])

# quý chưa kết thúc?
def q_end(y, q):
    return pd.Period(f"{y}Q{q}", freq="Q").end_time.date()

today = TODAY
in_progress = (any(q_end(sel_year, int(q[1:])) >= today for q in sel_q) if sel_q else sel_year == today.year)
partial_cmp = in_progress and len(sel_q) == 1

period_label = (f"Quý {' + '.join(ROMAN.get(int(q[1:]), q) for q in sel_q)} - {sel_year}" if sel_q else f"Năm {sel_year}")
logo = load_logo()
if logo["b64"]:
    img_tag = f"<img src='data:image/png;base64,{logo['b64']}' alt='Anna'>"
else:
    img_tag = f"<img src='{html.escape(logo['src'])}' referrerpolicy='no-referrer' alt='Anna'>"
# ảnh nền đặc hoặc logo sáng: đặt thẳng lên nền xanh; chỉ logo tối nền trong suốt mới cần ô trắng nhỏ
logo_html = img_tag if logo["plain"] else f"<div class='logo-tile'>{img_tag}</div>"
sub_txt = f"{period_label} · Dữ liệu đến {last_date:%d/%m/%Y}" + (" · kỳ báo cáo đang diễn ra" if in_progress else "")
st.markdown(
    f"<div class='hero'>{logo_html}<div><div class='eyebrow'>Anna · Bộ phận Pháp chế</div>"
    f"<div class='title'>BÁO CÁO CHẤT LƯỢNG CÔNG VIỆC PHÁP CHẾ - PHÒNG KIỂM SOÁT NỘI BỘ</div>"
    f"<div class='sub'>{sub_txt}</div></div>"
    f"<div class='pill'>{period_label} · {vn(len(f))} yêu cầu</div></div>", unsafe_allow_html=True)

if f.empty:
    st.info("Không có dữ liệu theo bộ lọc hiện tại.")
    st.stop()

S = stats(f)
SP = stats(prev_df) if prev_df is not None else None
PS = proc_stats(f)
PSP = proc_stats(prev_df) if prev_df is not None else None


def span_days(year, quarters_):
    # các ngày lịch thuộc kỳ; kỳ chưa kết thúc chỉ tính đến ngày có dữ liệu cuối
    qs = sorted(int(q[1:]) for q in quarters_) if quarters_ else [1, 2, 3, 4]
    days = set()
    for q_ in qs:
        per = pd.Period(f"{year}Q{q_}", freq="Q")
        a, b = per.start_time.normalize(), min(per.end_time.normalize(), last_date.normalize())
        if b >= a:
            days |= set(pd.date_range(a, b))
    return pd.DatetimeIndex(sorted(days))


def throughput(d, days_idx):
    # năng suất = số yêu cầu đã hoàn thành ÷ số ngày làm việc (T2-T6) trong kỳ
    nw = int((days_idx.dayofweek < 5).sum())
    done_ = int(d["is_done"].sum())
    return dict(done=done_, work_days=nw, cal_days=len(days_idx), per_day=(done_ / nw if nw else np.nan))


TP = throughput(f, span_days(sel_year, sel_q))
TPP = throughput(prev_df, span_days(*prev_key)) if prev_df is not None and prev_key else None

# ============================================================================
# PHẦN I
# ============================================================================
section("PHẦN I", "TỔNG QUAN CHẤT LƯỢNG CÔNG VIỆC PHÁP CHẾ",
        f"{period_label} · {vn(S['n'])} yêu cầu thuộc {S['cat']} hạng mục")

has_prev = SP is not None
cards = [
    kpi("Hạng mục", vn(S["cat"]), vn(SP["cat"]) if has_prev else None,
        delta_pill(S["cat"], SP["cat"] if has_prev else None, "neutral", "num"), "", "Số nhóm công việc", prev_label),
    kpi("Số lượng yêu cầu", vn(S["n"]), vn(SP["n"]) if has_prev else None,
        delta_pill(S["n"], SP["n"] if has_prev else None, "neutral", "pct", partial=partial_cmp), "", "Tổng yêu cầu tiếp nhận", prev_label),
    kpi("Đã hoàn thành", vn(S["done"]), vn(SP["done"]) if has_prev else None,
        delta_pill(S["done"], SP["done"] if has_prev else None, "neutral", "pct", partial=partial_cmp), "ok", "Trạng thái “Hoàn thành”", prev_label),
    kpi("Tỷ lệ hoàn thành", rate_txt(S["done"], S["n"]),
        rate_txt(SP["done"], SP["n"]) if has_prev else None,
        delta_pill(S["done"] / S["n"] * 100 if S["n"] else None,
                   SP["done"] / SP["n"] * 100 if has_prev and SP["n"] else None, "up_good", "pts"),
        "ok", f"{vn(S['done'])}/{vn(S['n'])} yêu cầu", prev_label),
    kpi("Đang thực hiện", vn(S["op"]), vn(SP["op"]) if has_prev else None,
        delta_pill(S["op"], SP["op"] if has_prev else None, "up_bad", "num", partial=False), "warn",
        "Chưa xử lý xong", prev_label),
]
if S["pa"] > 0:
    cards.append(kpi("Tạm dừng", vn(S["pa"]), vn(SP["pa"]) if has_prev else None,
                     delta_pill(S["pa"], SP["pa"] if has_prev else None, "up_bad", "num"), "pause",
                     "Yêu cầu bị tạm dừng", prev_label))
if S["oth"] > 0:
    cards.append(kpi("Trạng thái khác", vn(S["oth"]), None, "", "pause", "Chưa phân loại được"))
for col, html_ in zip(st.columns(len(cards), gap="medium"), cards):
    col.markdown(html_, unsafe_allow_html=True)

st.write("")
b1, b2, b3 = st.columns([1.6, 1, 1.35], gap="medium")

with b1:
    with card("cat"):
        ct("Cơ cấu theo hạng mục", "Số yêu cầu, số hoàn thành và tỷ lệ hoàn thành")
        t = (f.groupby("category").agg(n=("detail", "size"), d=("is_done", "sum"),
                                       op=("status", lambda s: int((s == ST_OPEN).sum())),
                                       pa=("status", lambda s: int((s == ST_PAUSE).sum())))
             .reset_index().sort_values("n", ascending=False))
        show_op, show_pa = t["op"].sum() > 0, t["pa"].sum() > 0
        h = "<table class='nt'><thead><tr><th>Hạng mục</th><th>Yêu cầu</th><th>Hoàn thành</th>"
        h += ("<th>Đang TH</th>" if show_op else "") + ("<th>Tạm dừng</th>" if show_pa else "") + "<th>Tỷ lệ hoàn thành</th></tr></thead><tbody>"
        for r in t.itertuples():
            h += f"<tr><td>{html.escape(r.category)}</td><td>{vn(r.n)}</td><td>{vn(r.d)}</td>"
            h += (num_cell(r.op) if show_op else "") + (num_cell(r.pa) if show_pa else "") + rate_cell(r.d, r.n) + "</tr>"
        h += f"<tr class='tot'><td>Tổng</td><td>{vn(S['n'])}</td><td>{vn(S['done'])}</td>"
        h += (f"<td>{vn(S['op'])}</td>" if show_op else "") + (f"<td>{vn(S['pa'])}</td>" if show_pa else "")
        h += rate_cell(S["done"], S["n"]) + "</tr></tbody></table>"
        st.markdown(h, unsafe_allow_html=True)

with b2:
    with card("quarter"):
        ct("Yêu cầu tiếp nhận theo quý", "Quý đang xem được tô đậm")
        tq = f_allq.groupby("quarter").size().reset_index(name="n").sort_values("quarter")
        sel_set = set(sel_q)
        fig = go.Figure(go.Bar(
            x=tq["quarter"], y=tq["n"], text=[vn(v) for v in tq["n"]], textposition="outside", cliponaxis=False,
            marker_color=[TEAL_DARK if (not sel_set or q in sel_set) else "#BFE3DE" for q in tq["quarter"]],
            width=0.55, hovertemplate="%{x}: %{y:,} yêu cầu<extra></extra>"))
        fig.update_yaxes(range=[0, max(tq["n"].max(), 1) * 1.18])
        show(style_fig(fig, 300))

with b3:
    with card("month"):
        ct("Khối lượng công việc theo thời gian", "Số yêu cầu tiếp nhận mỗi tháng · tháng thuộc kỳ đang xem được tô đậm")
        tm = f_allq.dropna(subset=["received"]).groupby("ym").size()
        if len(tm):
            idx = pd.period_range(tm.index.min(), tm.index.max(), freq="M").astype(str)
            tm = tm.reindex(idx, fill_value=0)
            labels = [f"T{p[5:7]}/{p[2:4]}" for p in idx]
            in_sel = [(not sel_set) or (f"Q{(int(p[5:7]) - 1) // 3 + 1}" in sel_set) for p in idx]
            vals = tm.values
            fig = go.Figure()
            fig.add_trace(go.Scatter(x=labels, y=vals, mode="lines", line=dict(color="#8CCBE8", width=2.5),
                                     fill="tozeroy", fillcolor="rgba(45,156,219,.10)", hoverinfo="skip"))
            fig.add_trace(go.Scatter(
                x=labels, y=vals, mode="markers+text", text=[vn(v) for v in vals], textposition="top center",
                cliponaxis=False, textfont=dict(size=11, color=[ORANGE if s else "#B7C3C1" for s in in_sel]),
                marker=dict(size=[10 if s else 6 for s in in_sel], color=[BLUE if s else "#B8D6E6" for s in in_sel]),
                hovertemplate="%{x}: %{y:,} yêu cầu<extra></extra>"))
            fig.update_xaxes(type="category")
            fig.update_yaxes(range=[0, max(vals.max(), 1) * 1.22])
            show(style_fig(fig, 300))

c1, c2, c3 = st.columns([1, 1.25, 1.25], gap="medium")

with c1:
    with card("status"):
        ct("Trạng thái thực hiện", "Cơ cấu theo trạng thái")
        order = [s for s in [ST_DONE, ST_OPEN, ST_PAUSE] if s in set(f["status"])] + \
                [s for s in f["status"].unique() if s not in (ST_DONE, ST_OPEN, ST_PAUSE)]
        ts = f["status"].value_counts().reindex(order)
        fig = go.Figure(go.Pie(labels=ts.index, values=ts.values, hole=0.7, sort=False, textinfo="none",
                               marker=dict(colors=[STATUS_COLORS.get(s, GREY) for s in ts.index],
                                           line=dict(color="#fff", width=2)),
                               hovertemplate="%{label}: %{value:,} (%{percent})<extra></extra>"))
        fig.add_annotation(text=f"<b style='font-size:26px'>{rate_txt(S['done'], S['n'])}</b><br>"
                                f"<span style='font-size:12px;color:{MUTED}'>hoàn thành</span>",
                           showarrow=False, x=0.5, y=0.5)
        show(style_fig(fig, 215))
        rows_ = "".join(
            f"<div class='legend-row'><span><i class='dot' style='background:{STATUS_COLORS.get(s, GREY)}'></i>{html.escape(s)}</span>"
            f"<span><b>{vn(n)}</b> <span class='n'>· {vn(n / S['n'] * 100, 1)}%</span></span></div>"
            for s, n in ts.items())
        st.markdown(rows_, unsafe_allow_html=True)

with c2:
    with card("unit"):
        ct("Đơn vị yêu cầu", "Số yêu cầu và tỷ trọng theo đơn vị")
        tu = f.groupby("unit").size().sort_values(ascending=False)
        if len(tu) > 8:
            tu = pd.concat([tu.head(7), pd.Series({f"Khác ({len(tu) - 7} đơn vị)": tu.iloc[7:].sum()})])
        tu = tu.sort_values(ascending=True)
        fig = go.Figure(go.Bar(
            y=tu.index, x=tu.values, orientation="h", marker_color=TEAL,
            text=[f"{vn(n)} · {vn(n / S['n'] * 100, 1)}%" for n in tu.values], textposition="outside", cliponaxis=False,
            hovertemplate="%{y}: %{x:,} yêu cầu<extra></extra>"))
        fig.update_xaxes(range=[0, tu.max() * 1.35], showticklabels=False, showgrid=False)
        fig.update_yaxes(showgrid=False)
        show(style_fig(fig, 340))

with c3:
    with card("concl"):
        box = st.empty()
        pk = f"{sel_year}|{'+'.join(sel_q) if sel_q else 'ALL'}"
        wk = f"note_{pk}"
        with st.expander("✍️ Nhập nhận xét bổ sung", expanded=False):
            note = st.text_area("Mỗi dòng là 1 ý", value=load_notes().get(pk, ""), key=wk, height=120,
                                placeholder="VD: Tháng 9 tăng đột biến do chương trình khuyến mại...",
                                on_change=save_note, args=(pk, wk))
            st.caption("Nhận xét được lưu theo từng kỳ báo cáo (file nhan_xet_bo_sung.json cạnh app).")

        top_cat = t.iloc[0]
        top2_unit = f.groupby("unit").size().sort_values(ascending=False).head(2)
        share_top2 = top2_unit.sum() / S["n"]
        b_ = [f"<b>Khối lượng công việc:</b> {period_label}, Phòng Pháp chế tiếp nhận <b>{vn(S['n'])}</b> yêu cầu thuộc <b>{S['cat']}</b> hạng mục."]
        if SP is not None and SP["n"]:
            if partial_cmp:
                b_[0] += f" Dữ liệu đến {last_date:%d/%m/%Y} (quý chưa kết thúc) nên chưa so sánh với {prev_label} ({vn(SP['n'])} yêu cầu)."
            else:
                ch = (S["n"] - SP["n"]) / SP["n"]
                b_[0] += f" So với {prev_label} ({vn(SP['n'])} yêu cầu): {'tăng' if ch >= 0 else 'giảm'} {vn(abs(ch) * 100, 1)}%."
        open_txt = []
        if S["op"]:
            open_txt.append(f"{vn(S['op'])} đang thực hiện")
        if S["pa"]:
            open_txt.append(f"{vn(S['pa'])} tạm dừng")
        if S["oth"]:
            open_txt.append(f"{vn(S['oth'])} trạng thái khác")
        b_.append(f"<b>Hiệu quả xử lý:</b> Tỷ lệ hoàn thành <b>{rate_txt(S['done'], S['n'])}</b> ({vn(S['done'])}/{vn(S['n'])} yêu cầu)"
                  + (f", năng suất xử lý trung bình {vn(TP['per_day'], 1)} yêu cầu/ngày làm việc." if TP["work_days"] else ".")
                  + (f" Còn {', '.join(open_txt)}." if open_txt else " Không còn yêu cầu tồn."))
        b_.append(f"<b>Cơ cấu công việc:</b> Hạng mục <b>{html.escape(top_cat['category'])}</b> chiếm tỷ trọng lớn nhất "
                  f"({vn(top_cat['n'])} yêu cầu, {vn(top_cat['n'] / S['n'] * 100, 1)}%).")
        b_.append("<b>Đơn vị yêu cầu:</b> " + " và ".join(f"<b>{html.escape(u)}</b>" for u in top2_unit.index)
                  + f" chiếm {vn(share_top2 * 100, 1)}% tổng số yêu cầu.")
        extra_lines = [re.sub(r"^[\-\•\*\s]+", "", ln).strip() for ln in (note or "").splitlines()]
        extra_lines = [ln for ln in extra_lines if ln]
        body = "<div class='concl'><h4>Kết luận quản trị:</h4><ul>" + "".join(f"<li>{x}</li>" for x in b_) + "</ul>"
        if extra_lines:
            body += "<div class='xt'>NHẬN XÉT BỔ SUNG</div><ul>" + "".join(f"<li>{html.escape(x)}</li>" for x in extra_lines) + "</ul>"
        box.markdown(body + "</div>", unsafe_allow_html=True)

d1, d2 = st.columns([1.5, 1], gap="medium")
with d1:
    with card("heat"):
        ct("Đơn vị yêu cầu × Hạng mục", "Màu càng đậm = càng nhiều yêu cầu")
        hm = f.pivot_table(index="unit", columns="category", values="detail", aggfunc="size", fill_value=0)
        hm = hm.loc[hm.sum(axis=1).sort_values().index, hm.sum(axis=0).sort_values(ascending=False).index]
        fig = go.Figure(go.Heatmap(
            z=hm.values, x=["<br>".join(textwrap.wrap(str(c_), 11)) for c_ in hm.columns], y=list(hm.index), xgap=3, ygap=3, showscale=False,
            text=[[vn(v) if v else "" for v in row] for row in hm.values], texttemplate="%{text}",
            colorscale=[[0, "#F0F9F7"], [1, TEAL_DARK]], hovertemplate="%{y} · %{x}: %{z:,}<extra></extra>"))
        fig.update_xaxes(side="top", tickangle=0, tickfont=dict(size=11))
        show(style_fig(fig, 350, top=62))
with d2:
    with card("wday"):
        ct("Yêu cầu theo thứ trong tuần", "Ngày nhận yêu cầu")
        wd = f["received"].dt.dayofweek.value_counts().reindex(range(7), fill_value=0)
        fig = go.Figure(go.Bar(x=["T2", "T3", "T4", "T5", "T6", "T7", "CN"], y=wd.values, marker_color=TEAL,
                               text=[vn(v) if v else "" for v in wd.values], textposition="outside", cliponaxis=False,
                               width=0.6))
        fig.update_yaxes(range=[0, max(wd.max(), 1) * 1.18])
        show(style_fig(fig, 330))

st.markdown(
    "<div class='fn'><b>Trong đó:</b><br>"
    "- <b>Đã hoàn thành</b>: số yêu cầu có trạng thái “Hoàn thành”. <b>Đang thực hiện / Tạm dừng</b>: theo cột Trạng thái.<br>"
    "- <b>Tỷ lệ hoàn thành</b> = Đã hoàn thành / Số lượng yêu cầu, làm tròn xuống 1 số lẻ (chỉ hiển thị 100% khi toàn bộ đã hoàn thành).<br>"
    "- Kỳ so sánh chỉ hiển thị khi chọn đúng 1 quý; các bộ lọc khác (đơn vị, trạng thái, người thực hiện, hạng mục) được áp dụng cho cả 2 kỳ.</div>",
    unsafe_allow_html=True)

# ============================================================================
# PHẦN II
# ============================================================================
section("PHẦN II", "TỔNG QUAN BỘ PHẬN PHÁP CHẾ",
        "Khối lượng, tiến độ và năng suất theo từng cán bộ · năng suất = số yêu cầu hoàn thành ÷ số ngày làm việc")

pcards = [
    kpi("Năng suất xử lý TB", f"{vn(TP['per_day'], 1)} <small>yêu cầu/ngày</small>" if TP["work_days"] else "–",
        f"{vn(TPP['per_day'], 1)} yêu cầu/ngày" if TPP and TPP["work_days"] else None,
        delta_pill(TP["per_day"], TPP["per_day"] if TPP and TPP["work_days"] else None, "up_good", "dec", " yêu cầu"), "",
        f"{vn(TP['done'])} hoàn thành ÷ {vn(TP['work_days'])} ngày làm việc", prev_label),
    kpi("Hoàn thành trong ngày", rate_txt(PS["same"], PS["n"]) if PS["n"] else "–",
        rate_txt(PSP["same"], PSP["n"]) if PSP and PSP["n"] else None,
        delta_pill(PS["same"] / PS["n"] * 100 if PS["n"] else None,
                   PSP["same"] / PSP["n"] * 100 if PSP and PSP["n"] else None, "up_good", "pts"), "ok",
        f"{vn(PS['same'])}/{vn(PS['n'])} yêu cầu", prev_label),
    kpi("Xử lý trên 7 ngày", share_txt(PS["slow"], PS["n"]) if PS["n"] else "–",
        share_txt(PSP["slow"], PSP["n"]) if PSP and PSP["n"] else None,
        delta_pill(PS["slow"] / PS["n"] * 100 if PS["n"] else None,
                   PSP["slow"] / PSP["n"] * 100 if PSP and PSP["n"] else None, "up_bad", "pts"), "warn",
        f"{vn(PS['slow'])}/{vn(PS['n'])} yêu cầu", prev_label),
]
for col, html_ in zip(st.columns(3, gap="medium"), pcards):
    col.markdown(html_, unsafe_allow_html=True)
st.write("")

people = f["owner"].value_counts().index.tolist()
pcol = {p: PERSON_COLORS[i % len(PERSON_COLORS)] for i, p in enumerate(people)}
cat_order = f.groupby("category").size().sort_values(ascending=False).index.tolist()

e1, e2 = st.columns([1.6, 1], gap="medium")
with e1:
    with card("pivot"):
        ct("Hạng mục × Người thực hiện", "Số yêu cầu và tỷ lệ hoàn thành của từng cán bộ")
        st.markdown(pivot_html(f, people, cat_order, "Hạng mục"), unsafe_allow_html=True)
with e2:
    with card("pq"):
        ct("Yêu cầu tiếp nhận theo quý", "Theo cán bộ phụ trách")
        tqp = f_allq.groupby(["quarter", "owner"]).size().reset_index(name="n")
        fig = go.Figure()
        for p in people:
            s = tqp[tqp["owner"] == p]
            fig.add_trace(go.Bar(x=s["quarter"], y=s["n"], name=p, marker_color=pcol[p],
                                 text=[vn(v) for v in s["n"]], textposition="inside", width=0.55,
                                 hovertemplate=f"{p} · %{{x}}: %{{y:,}}<extra></extra>"))
        fig.update_layout(barmode="stack")
        fig.update_xaxes(categoryorder="category ascending")
        show(style_fig(fig, 300, legend=True))

g1, g2 = st.columns([1, 1.4], gap="medium")
with g1:
    with card("pstatus"):
        ct("Tỷ lệ xử lý yêu cầu", "Cơ cấu trạng thái của từng cán bộ")
        ps = f.groupby(["owner", "status"]).size().reset_index(name="n")
        ps["pct"] = ps["n"] / ps.groupby("owner")["n"].transform("sum") * 100
        tot_by = f.groupby("owner").size()
        fig = go.Figure()
        for s in order:
            sub = ps[ps["status"] == s].set_index("owner").reindex(people[::-1])
            txt = [(f"{vn(r.pct, 1)}%" if (not pd.isna(r.pct) and (r.pct >= 8 or s != ST_DONE and r.pct > 0)) else "")
                   for r in sub.itertuples()]
            if s == ST_DONE:
                txt = [rate_txt(int(sub.loc[o, "n"]) if not pd.isna(sub.loc[o, "n"]) else 0, int(tot_by[o])) for o in sub.index]
            fig.add_trace(go.Bar(y=sub.index, x=sub["pct"].fillna(0), name=s, orientation="h",
                                 marker_color=STATUS_COLORS.get(s, GREY), text=txt, textposition="inside",
                                 insidetextanchor="middle", textfont=dict(color="#fff", size=11),
                                 hovertemplate=f"{s}: %{{x:.2f}}%<extra></extra>"))
        fig.update_layout(barmode="stack")
        fig.update_xaxes(range=[0, 100], tickvals=[0, 25, 50, 75, 100], ticktext=["0%", "25%", "50%", "75%", "100%"])
        show(style_fig(fig, 300, legend=True).update_layout(margin=dict(l=6, r=26, t=34, b=6)))
with g2:
    with card("pmonth"):
        ct("Khối lượng công việc theo thời gian", "Số yêu cầu tiếp nhận mỗi tháng theo cán bộ")
        tmp = f_allq.dropna(subset=["received"]).groupby(["ym", "owner"]).size().reset_index(name="n")
        if len(tmp):
            idx = pd.period_range(tmp["ym"].min(), tmp["ym"].max(), freq="M").astype(str)
            labels = {p: f"T{p[5:7]}/{p[2:4]}" for p in idx}
            fig = go.Figure()
            for p in people:
                s = tmp[tmp["owner"] == p].set_index("ym")["n"].reindex(idx, fill_value=0)
                fig.add_trace(go.Scatter(x=[labels[i] for i in idx], y=s.values, name=p, mode="lines+markers+text",
                                         line=dict(color=pcol[p], width=2.5), marker=dict(size=7),
                                         text=[vn(v) for v in s.values], textposition="top center",
                                         textfont=dict(size=10, color=pcol[p]), cliponaxis=False,
                                         hovertemplate=f"{p} · %{{x}}: %{{y:,}}<extra></extra>"))
            fig.update_xaxes(type="category")
            fig.update_yaxes(range=[0, tmp.groupby(["ym", "owner"])["n"].sum().max() * 1.2])
            show(style_fig(fig, 300, legend=True))

h1, h2 = st.columns([1.6, 1], gap="medium")
with h1:
    with card("perf"):
        ct("Hiệu suất từng cán bộ", "Yêu cầu/ngày = hoàn thành ÷ ngày làm việc · Trong ngày và Trên 7 ngày tính trên các yêu cầu đã hoàn thành")
        show_pa_p = bool(f["status"].eq(ST_PAUSE).any())
        pa_total_cell = f"<td>{vn(S['pa'])}</td>" if show_pa_p else ""
        h = ("<table class='nt'><thead><tr><th>Cán bộ</th><th>Yêu cầu</th><th>Hoàn thành</th><th>Tỷ lệ HT</th>"
             "<th>Đang TH</th>" + ("<th>Tạm dừng</th>" if show_pa_p else "") +
             "<th>Yêu cầu/ngày</th><th>Trong ngày</th><th>Trên 7 ngày</th></tr></thead><tbody>")
        for p in people:
            sub = f[f["owner"] == p]
            ss, pp = stats(sub), proc_stats(sub)
            h += (f"<tr><td><i class='dot' style='background:{pcol[p]}'></i>{html.escape(p)}</td><td>{vn(ss['n'])}</td>"
                  f"<td>{vn(ss['done'])}</td><td class='{'low' if ss['done'] < ss['n'] else ''}'>{rate_txt(ss['done'], ss['n'])}</td>"
                  f"{num_cell(ss['op'])}{num_cell(ss['pa']) if show_pa_p else ''}"
                  f"<td>{vn(ss['done'] / TP['work_days'], 1) if TP['work_days'] else '–'}</td>"
                  f"<td>{rate_txt(pp['same'], pp['n']) if pp['n'] else '–'}</td>"
                  f"<td>{share_txt(pp['slow'], pp['n']) if pp['n'] else '–'}</td></tr>")
        h += (f"<tr class='tot'><td>Tổng</td><td>{vn(S['n'])}</td><td>{vn(S['done'])}</td><td>{rate_txt(S['done'], S['n'])}</td>"
              f"<td>{vn(S['op'])}</td>{pa_total_cell}"
              f"<td>{vn(TP['per_day'], 1) if TP['work_days'] else '–'}</td>"
              f"<td>{rate_txt(PS['same'], PS['n']) if PS['n'] else '–'}</td>"
              f"<td>{share_txt(PS['slow'], PS['n']) if PS['n'] else '–'}</td></tr></tbody></table>")
        st.markdown(h, unsafe_allow_html=True)
with h2:
    with card("bucket"):
        ct("Phân bố thời gian xử lý", "% yêu cầu đã hoàn thành theo mốc thời gian · từng cán bộ")
        dv = f[f["days"].notna()].copy()
        if len(dv):
            dv["bucket"] = pd.cut(dv["days"], [-1, 0, 2, 7, 10 ** 6], labels=["Trong ngày", "1–2 ngày", "3–7 ngày", "Trên 7 ngày"])
            bc = dv.groupby(["owner", "bucket"], observed=False).size().reset_index(name="n")
            bc["pct"] = bc["n"] / bc.groupby("owner")["n"].transform("sum") * 100
            fig = go.Figure()
            for b, colr in zip(["Trong ngày", "1–2 ngày", "3–7 ngày", "Trên 7 ngày"], [TEAL_DARK, "#7ACBC3", ORANGE, "#C2501F"]):
                s = bc[bc["bucket"] == b].set_index("owner").reindex(people[::-1])
                fig.add_trace(go.Bar(y=s.index, x=s["pct"].fillna(0), name=b, orientation="h", marker_color=colr,
                                     text=[f"{vn(v, 0)}%" if v >= 6 else "" for v in s["pct"].fillna(0)],
                                     textposition="inside", textfont=dict(color="#fff", size=11),
                                     hovertemplate=f"{b}: %{{x:.1f}}%<extra></extra>"))
            fig.update_layout(barmode="stack")
            fig.update_xaxes(range=[0, 100], tickvals=[0, 25, 50, 75, 100], ticktext=["0%", "25%", "50%", "75%", "100%"])
            show(style_fig(fig, 250, legend=True).update_layout(margin=dict(l=6, r=26, t=34, b=6)))
        else:
            st.caption("Chưa có yêu cầu hoàn thành có đủ ngày nhận và ngày hoàn thành.")

st.markdown(
    "<div class='fn'><b>Trong đó:</b><br>"
    "- <b>Năng suất xử lý TB</b> = số yêu cầu đã hoàn thành / số ngày làm việc (thứ Hai–thứ Sáu) trong kỳ; kỳ chưa kết thúc chỉ tính đến ngày có dữ liệu cuối, chưa trừ ngày lễ.<br>"
    "- <b>Thời gian xử lý</b> của mỗi yêu cầu = Ngày hoàn thành − Ngày nhận yêu cầu (theo ngày lịch, không tính giờ); chỉ tính các yêu cầu <b>đã hoàn thành</b> và có đủ 2 ngày hợp lệ.<br>"
    "- <b>Hoàn thành trong ngày</b> = số yêu cầu xử lý 0 ngày / số yêu cầu đã hoàn thành có thời gian xử lý hợp lệ.<br>"
    "- <b>Xử lý trên 7 ngày</b> = số yêu cầu hoàn thành sau hơn 7 ngày / số yêu cầu đã hoàn thành có thời gian xử lý hợp lệ.</div>",
    unsafe_allow_html=True)

with st.expander("🔍 Đối chiếu năng suất và thời gian xử lý"):
    st.markdown(
        "**Năng suất xử lý TB** = số yêu cầu đã hoàn thành ÷ số ngày làm việc (thứ Hai–thứ Sáu) trong kỳ. "
        "Kỳ chưa kết thúc chỉ tính đến ngày có dữ liệu cuối; chưa trừ ngày lễ.")
    ref_ = pd.DataFrame({
        "Chỉ tiêu": ["Số yêu cầu đã hoàn thành", "Số ngày lịch trong kỳ", "Số ngày làm việc (T2–T6)",
                     "Năng suất = hoàn thành ÷ ngày làm việc", "Tham khảo: hoàn thành ÷ ngày lịch"],
        "Giá trị": [vn(TP["done"]), vn(TP["cal_days"]), vn(TP["work_days"]),
                    f"{vn(TP['per_day'], 2)} yêu cầu/ngày" if TP["work_days"] else "–",
                    f"{vn(TP['done'] / TP['cal_days'], 2)} yêu cầu/ngày" if TP["cal_days"] else "–"]})
    st.dataframe(ref_, hide_index=True, width="stretch")
    dv_ = f[f["days"].notna()]
    if len(dv_):
        st.markdown("**Phân bố thời gian xử lý** (số ngày lịch từ Ngày nhận đến Ngày hoàn thành, cùng ngày = 0) "
                    "· dùng cho tỷ lệ “hoàn thành trong ngày” và “xử lý trên 7 ngày”:")
        bins_ = dv_["days"].clip(upper=8).astype(int)
        tb = pd.DataFrame({"Số yêu cầu": bins_.value_counts().reindex(range(9), fill_value=0)})
        tb.index = ["0 (trong ngày)"] + [str(i) for i in range(1, 8)] + ["Trên 7 ngày"]
        tb["Tỷ trọng"] = [f"{vn(v / len(dv_) * 100, 1)}%" for v in tb["Số yêu cầu"]]
        st.dataframe(tb.reset_index(names="Số ngày xử lý"), hide_index=True, width="stretch")
    n_excl = int(f["is_done"].sum()) - len(dv_)
    if n_excl > 0:
        st.caption(f"Có {vn(n_excl)} yêu cầu đã hoàn thành nhưng không tính được thời gian xử lý "
                   f"(thiếu Ngày hoàn thành hoặc Ngày hoàn thành sớm hơn Ngày nhận) nên không nằm trong phân bố trên.")

# ============================================================================
# PHẦN III
# ============================================================================
section("PHẦN III", "CHI TIẾT CÔNG VIỆC", "Danh sách theo bộ lọc bên trái · có thể tải về file Excel để theo dõi")

only_open = st.toggle("Chỉ xem các yêu cầu chưa hoàn thành", value=False)
fd = f[~f["is_done"]] if only_open else f
detail = (fd.sort_values(["received", "detail"])[
    ["detail", "status", "received", "done", "days", "owner", "unit", "category", "note"]]
    .rename(columns={"detail": "Chi Tiết Yêu Cầu", "status": "Trạng Thái", "received": "Ngày nhận yêu cầu",
                     "done": "Ngày hoàn thành", "days": "Số ngày xử lý", "owner": "Cán bộ phụ trách",
                     "unit": "Đơn vị yêu cầu", "category": "Hạng mục", "note": "Ghi chú"}))

t1, t2 = st.columns([4, 1])
t1.caption(f"Hiển thị {vn(len(detail))} yêu cầu.")
t2.download_button("⬇️ Tải file Excel", data=to_excel(detail),
                   file_name=f"chi_tiet_cong_viec_phap_che_{sel_year}_{TODAY:%Y%m%d}.xlsx",
                   mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                   width="stretch", type="primary")


def color_status(v):
    return {ST_DONE: "color:#0B7F75;font-weight:600", ST_OPEN: "color:#C2641B;font-weight:600",
            ST_PAUSE: "color:#27348B;font-weight:600"}.get(v, "")


st.dataframe(
    detail.style.map(color_status, subset=["Trạng Thái"]),
    width="stretch", height=540, hide_index=True,
    column_config={
        "Chi Tiết Yêu Cầu": st.column_config.TextColumn(width="large"),
        "Ngày nhận yêu cầu": st.column_config.DateColumn(format="DD/MM/YYYY"),
        "Ngày hoàn thành": st.column_config.DateColumn(format="DD/MM/YYYY"),
        "Số ngày xử lý": st.column_config.NumberColumn(format="%d"),
    },
)

# ---------- Kiểm tra chất lượng dữ liệu ----------
with st.expander("🔍 Kiểm tra chất lượng dữ liệu nguồn (theo bộ lọc hiện tại)"):
    checks = {
        "Trạng thái “Hoàn thành” nhưng thiếu Ngày hoàn thành (không tính được thời gian xử lý)": f["dq_done_nodate"],
        "Có Ngày hoàn thành nhưng trạng thái chưa là “Hoàn thành”": f["dq_date_notdone"],
        "Ngày hoàn thành sớm hơn Ngày nhận yêu cầu": f["dq_negative"],
        "Trạng thái ngoài 3 giá trị chuẩn (Hoàn thành / Đang thực hiện / Tạm dừng)": f["dq_other_status"],
    }
    dq = pd.DataFrame({"Kiểm tra": list(checks), "Số dòng": [int(v.sum()) for v in checks.values()]})
    st.dataframe(dq, hide_index=True, width="stretch")
    if n_no_received:
        st.caption(f"⚠️ Có {vn(n_no_received)} dòng trong nguồn thiếu Ngày nhận yêu cầu nên không nằm trong báo cáo.")
    bad = [k for k, v in checks.items() if v.sum() > 0]
    if bad:
        pick = st.selectbox("Xem các dòng bị cảnh báo", bad)
        st.dataframe(f.loc[checks[pick], ["detail", "status", "received", "done", "owner"]]
                     .rename(columns={"detail": "Chi Tiết Yêu Cầu", "status": "Trạng Thái", "received": "Ngày nhận",
                                      "done": "Ngày hoàn thành", "owner": "Cán bộ phụ trách"}),
                     hide_index=True, width="stretch")
    else:
        st.success("Không phát hiện bất thường trong dữ liệu đang lọc.")

st.caption("Báo cáo tự cập nhật theo dữ liệu nguồn.")
